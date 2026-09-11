#!/usr/bin/env python3
"""Converte a planilha HEF (Tipo / Marca / Numero Serie / Local) no JSON oficial.

TAG = HEF-0001, HEF-0002… na ordem da planilha (após dedup tipo+série).
O número de série da coluna vai para nSerie — nunca para a TAG.
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import openpyxl

MULTI_FAB = [
    "Nihon Kohden",
    "Konica Minolta",
    "Nova Instruments",
    "Medtronic",
    "Alfamed",
    "Mindray",
    "Medcaptain",
    "HillRom",
    "EDAN",
    "Fanem",
    "Daquino",
    "WEM",
    "Fresenius",
    "Saubern",
    "Barrfab",
    "Welmy",
    "Ortosintese",
    "Olympus",
    "Shimadzu",
    "Medrad",
    "OQTIS",
]

PLANO = {
    "monitor multiparametrico": "Monitor",
    "monitor fisiológico ecg": "Monitor",
    "monitor contraste": "Monitor",
    "monitor torre video": "Monitor",
    "ventilador pulmonar": "Ventilador",
    "ventilador": "Ventilador",
    "cpap": "Ventilador",
    "bomba infusão": "Bomba de Infusão",
    "bomba seringa": "Bomba de Infusão",
    "bomba enteral": "Bomba de Infusão",
    "cardioversor": "Desfibrilador",
    "eletrocardiografo": "Eletrocardiógrafo",
    "raio x": "Raio-X",
    "raio x portátil": "Raio-X",
    "ultrassom": "Ultrassom",
    "autoclave universal": "Autoclave",
    "autoclave baixa temperatura": "Autoclave",
}


def norm_serie(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".")[0]
    return s


def split_marca(marca: str) -> tuple[str, str]:
    m = re.sub(r"\s+", " ", marca.replace("–", "-").replace("—", "-")).strip()
    for fab in sorted(MULTI_FAB, key=len, reverse=True):
        if m.lower().startswith(fab.lower()):
            rest = m[len(fab) :].strip(" -")
            return ("EDAN" if fab == "EDAN" else fab), (rest or "Não informado")
    parts = m.split(" ", 1)
    if len(parts) == 1:
        return parts[0], "Não informado"
    return parts[0], parts[1]


def storage_score(setor: str) -> int:
    s = setor.lower()
    if "depósito" in s or "deposito" in s:
        return 2
    if "sala equipamentos" in s:
        return 1
    return 0


def convert(xlsx: Path) -> dict:
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for i, row in enumerate(ws.iter_rows(min_row=3, max_row=ws.max_row, max_col=4, values_only=True), start=3):
        tipo, marca, serie, local = row
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        tipo = str(tipo).strip()
        marca = str(marca).strip()
        local = str(local).strip()
        serie = norm_serie(serie)
        fab, modelo = split_marca(marca)
        rows.append(
            {
                "linha": i,
                "tipo": tipo,
                "fab": fab,
                "modelo": modelo,
                "serie": serie,
                "setor": local,
                "plano": PLANO.get(tipo.lower(), tipo),
            }
        )

    best: dict[tuple[str, str], dict] = {}
    order: list[dict] = []
    for r in rows:
        if not r["serie"]:
            order.append(r)
            continue
        key = (r["tipo"].lower(), r["serie"].lower())
        if key not in best:
            best[key] = r
            order.append(r)
            continue
        old = best[key]
        if storage_score(r["setor"]) < storage_score(old["setor"]) or (
            storage_score(r["setor"]) == storage_score(old["setor"]) and r["linha"] > old["linha"]
        ):
            for i, x in enumerate(order):
                if x is old:
                    order[i] = r
                    break
            best[key] = r

    avisos: list[str] = []
    for r in rows:
        if not r["serie"]:
            continue
        winner = best[(r["tipo"].lower(), r["serie"].lower())]
        if winner is not r:
            avisos.append(
                f"Duplicata {r['tipo']} série {r['serie']}: linha {r['linha']} ({r['setor']}) "
                f"descartada; mantida linha {winner['linha']} ({winner['setor']})."
            )

    serie_tipos: dict[str, list[str]] = {}
    for r in order:
        if r["serie"]:
            serie_tipos.setdefault(r["serie"].lower(), []).append(r["tipo"])
    for serie, tipos in serie_tipos.items():
        if len(tipos) > 1:
            avisos.append(
                f"Série {serie} compartilhada por tipos distintos ({', '.join(tipos)}) — "
                "nSerie repetido (schema permite); TAGs HEF-NNNN distintas."
            )

    eqs = []
    for i, r in enumerate(order, start=1):
        tag = f"HEF-{i:04d}"
        if not r["serie"]:
            avisos.append(
                f"Linha {r['linha']}: sem nº de série — nSerie vazio; tag {tag} "
                f"({r['tipo']} / {r['setor']})"
            )
        eqs.append(
            {
                "tag": tag,
                "nome": r["tipo"],
                "planoDescricao": r["plano"],
                "fabricante": r["fab"],
                "modelo": r["modelo"],
                "setor": r["setor"],
                "patrimonio": None,
                "nSerie": r["serie"],
                "registroAnvisa": None,
                "validadeAnvisa": None,
                "dataAquisicao": None,
                "dataInstalacao": None,
                "valorAquisicao": None,
                "situacao": "ATIVO",
                "observacao": None,
                "laudos": [],
            }
        )

    return {
        "meta": {
            "fonte": f"{xlsx.name} — Levantamento de Equipamentos Biomédicos (aba Planilha1)",
            "totalLinhas": len(rows),
            "totalEquipamentos": len(eqs),
            "formatoTag": "HEF-NNNN (4 dígitos, HEF-0001…)",
            "ordemTags": (
                "Sequência HEF-NNNN na ordem da planilha original (linhas 3+), "
                "após deduplicar tipo+série: permanece o registro de menor score de "
                "depósito/sala de equipamentos ou, empate, a maior linha. "
                "O nº de série da coluna fica em nSerie; TAG não reutiliza a série."
            ),
            "avisos": avisos,
        },
        "equipamentos": eqs,
    }


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("xlsx")
    p.add_argument("-o", "--output", default="scripts/dados/equipamentos-reais.json")
    args = p.parse_args()
    payload = convert(Path(args.xlsx))
    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(out), **payload["meta"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
